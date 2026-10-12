"""Exercise retained upstream old dependencies under the genuine native init."""
import io
import json
from pathlib import Path
import tempfile
import app
from sqlalchemy_imageattach.stores.fs import HttpExposedFileSystemStore
from docx import Document

with tempfile.TemporaryDirectory(prefix='iris-upgrade-compat-') as directory:
    root = Path(directory)
    store = HttpExposedFileSystemStore(str(root / 'images'))
    binary = b'IRIS image store compatibility\x00\xff'
    store.put_file(io.BytesIO(binary), 'acceptance', 42, 1, 1, 'image/png', False)
    with store.get_file('acceptance', 42, 1, 1, 'image/png') as stream:
        assert stream.read() == binary
    document = Document()
    document.add_paragraph('IRIS Python 3.12 retained document compatibility')
    path = root / 'acceptance.docx'
    document.save(str(path))
    restored = Document(str(path))
    assert restored.paragraphs[0].text == 'IRIS Python 3.12 retained document compatibility'
print(json.dumps({'image_store_write_read': True, 'python_docx_save_reopen': True}))
