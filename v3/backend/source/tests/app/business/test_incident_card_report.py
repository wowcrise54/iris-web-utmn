"""Both exports must render the same snapshot, escape HTML, and retain long text."""

import base64
import importlib.util
from io import BytesIO
from pathlib import Path
import unittest

from docx import Document


MODULE = Path(__file__).parents[3] / "app/business/reports/incident_card_report.py"


class IncidentCardReportTests(unittest.TestCase):
    def render(self, snapshot, **kwargs):
        self.assertTrue(
            MODULE.is_file(), "Incident card report renderer is not implemented"
        )
        spec = importlib.util.spec_from_file_location("incident_card_report", MODULE)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module.build_export(snapshot, **kwargs)

    def snapshot(self):
        return {
            "revision": 7,
            "meta": {"case_id": 2},
            "warnings": [],
            "fields": {
                "description": {
                    "value": "<script>alert(1)</script> Кириллица " + "текст " * 1500
                },
                "detected_at": {"value": "2026-10-13T06:00:00+00:00"},
                "categories": {"value": [2, 3]},
                "completed": {"value": False},
                "acquainted": {
                    "value": [{"name": "Иванов Иван", "position": "Аналитик"}]
                },
            },
            "derived": {"importance": "high", "priority": 1, "reaction_seconds": 90},
            "catalog": {
                "sections": [
                    {
                        "id": "details",
                        "label_ru": "Детали инцидента",
                        "fields": [
                            {
                                "key": "description",
                                "label_ru": "Описание",
                                "type": "textarea",
                            },
                            {
                                "key": "detected_at",
                                "label_ru": "Обнаружение",
                                "type": "datetime",
                            },
                            {
                                "key": "categories",
                                "label_ru": "Категории",
                                "type": "multichoice",
                                "options": [
                                    {"value": 2, "label_ru": "Вредоносное ПО"},
                                    {"value": 3, "label_ru": "Доступ"},
                                ],
                            },
                            {
                                "key": "completed",
                                "label_ru": "Завершён",
                                "type": "boolean",
                            },
                            {
                                "key": "acquainted",
                                "label_ru": "Ознакомлены",
                                "type": "people",
                            },
                        ],
                    }
                ]
            },
        }

    def test_same_snapshot_safe_html_and_complete_docx(self):
        snapshot = self.snapshot()
        result = self.render(snapshot)
        self.assertEqual(result["revision"], 7)
        self.assertEqual(result["filename"], "incident-card-2.docx")
        self.assertNotIn("<script>", result["html"])
        self.assertIn("&lt;script&gt;", result["html"])
        self.assertIn("13.10.2026 11:00", result["html"])
        doc = Document(BytesIO(base64.b64decode(result["docx_base64"])))
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn(snapshot["fields"]["description"]["value"], text)
        for expected in ["Вредоносное ПО", "Доступ", "Иванов Иван", "Аналитик", "Нет"]:
            self.assertIn(expected, text)
            self.assertIn(expected, result["html"])

    def test_empty_snapshot_retains_form_and_blank_signature_lines(self):
        snapshot = self.snapshot()
        snapshot["fields"] = {}
        snapshot["derived"] = {}
        result = self.render(snapshot)
        self.assertIn("Приложение № 4", result["html"])
        self.assertIn("Подпись", result["html"])
        self.assertIn("Не определено", result["html"])
        self.assertIn("@media print", result["html"])

    def test_explicit_timezone_is_used_in_both_exports(self):
        result = self.render(self.snapshot(), timezone_name="UTC")
        self.assertIn("13.10.2026 06:00", result["html"])
        doc = Document(BytesIO(base64.b64decode(result["docx_base64"])))
        text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn("13.10.2026 06:00", text)

    def test_priority_tables_precede_investigation_and_signatures_use_people(self):
        snapshot = self.snapshot()
        catalog_spec = importlib.util.spec_from_file_location(
            "real_card_catalog", MODULE.parent.parent / "incident_card_catalog.py"
        )
        catalog = importlib.util.module_from_spec(catalog_spec)
        catalog_spec.loader.exec_module(catalog)
        snapshot["catalog"]["sections"] = catalog.SECTIONS
        snapshot["fields"]["prepared_by"] = {"value": "Петров Пётр"}
        snapshot["fields"]["prepared_position"] = {"value": "Специалист ИБ"}
        result = self.render(snapshot)
        self.assertLess(
            result["html"].index("Важность по таблице регламента"),
            result["html"].index("<h2>Расследование"),
        )
        self.assertIn("Петров Пётр", result["html"])
        self.assertIn("Иванов Иван", result["html"])
        doc = Document(BytesIO(base64.b64decode(result["docx_base64"])))
        tables = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn("Специалист ИБ", tables)


if __name__ == "__main__":
    unittest.main()
