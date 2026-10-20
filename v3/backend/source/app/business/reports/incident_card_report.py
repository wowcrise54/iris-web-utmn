"""Native Appendix 4 form: editable Word elements and print from the same document."""

import base64
from datetime import datetime
from html import escape
from io import BytesIO
from pathlib import Path
import re
from zoneinfo import ZoneInfo

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import ImageFont

from app.business.incident_card_catalog import FIELDS

IMPORTANCE = {"low": "Низкая", "medium": "Средняя", "high": "Высокая"}


def _clean(value):
    return (
        re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))
        .replace("\r\n", "\n")
        .replace("\r", "\n")
    )


def _font(size):
    for name in (
        "C:/Windows/Fonts/times.ttf",
        "/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSerif-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
    ):
        if Path(name).exists():
            return ImageFont.truetype(name, size)
    return ImageFont.load_default(size=size)


def _lines(text, width, font):
    # Measure actual glyph widths; never clip even unbroken URLs or hashes.
    lines = []
    for paragraph in text.split("\n"):
        current = ""
        for char in paragraph:
            if current and font.getlength(current + char) > width:
                lines.append(current)
                current = ""
            current += char
        lines.append(current)
    return lines


def _values(snapshot, zone):
    values = {}
    for key, field in snapshot.get("fields", {}).items():
        value = field.get("value")
        if value is None:
            values[key] = ""
        elif key.endswith("_at"):
            try:
                value = (
                    datetime.fromisoformat(value.replace("Z", "+00:00"))
                    .astimezone(ZoneInfo(zone))
                    .strftime("%d.%m.%Y %H:%M")
                )
            except (ValueError, TypeError):
                pass
            values[key] = _clean(value)
        else:
            values[key] = (
                _clean(value) if not isinstance(value, (list, bool)) else value
            )
    preparers = values.get("prepared_by")
    if isinstance(preparers, list):
        first = preparers[0] if preparers else {}
        values["prepared_by"] = _clean(first.get("name") or "")
        values["prepared_position"] = _clean(first.get("position") or "")
    if values.get("unacceptable_event") is not True:
        values["unacceptable_event_kinds"] = []
    derived = snapshot.get("derived", {})
    values["importance"] = IMPORTANCE.get(derived.get("importance"), "")
    values["priority"] = str(derived.get("priority") or "")
    seconds = derived.get("incident_duration_seconds")
    if seconds is not None:
        days, rest = divmod(int(seconds), 86400)
        hours, rest = divmod(rest, 3600)
        values["duration"] = f"{days} дн. {hours} ч. {rest // 60} мин."
    people = values.get("acquainted") or []
    for i, person in enumerate(people[:2]):
        for key in ("name", "position"):
            values[f"person_{i}_{key}"] = _clean(person.get(key) or "")
    return values


TEMPLATE = Path(__file__).with_name("incident_card_template") / "appendix4.docx"

# Slots in the supplied, versioned DOCX. The source template owns all formatting.
INLINE_FIELDS = {
    2: "id",
    3: "incident_at",
    4: "detection_source",
    5: "criticality",
    7: "reporter_name",
    8: "reporter_phone",
    9: "reporter_email",
    10: "incident_location",
    11: "reporter_organization",
    17: "affected_assets",
    19: "occurred_at",
    20: "detected_at",
    21: "notified_at",
    32: "duration",
    39: "vulnerabilities",
    49: "asset_data",
    50: "asset_hardware",
    51: "asset_software",
    52: "asset_communications",
    53: "asset_documentation",
    61: "importance",
    64: "priority",
    67: "investigation_started_at",
    68: "investigators",
    69: "incident_ended_at",
    70: "impact_ended_at",
    71: "investigation_ended_at",
    72: "report_url",
}
NARRATIVE_FIELDS = {
    "description": [14, 15, 16],
    "business_impact": [55, 56, 57, 58],
    "offender_description": [75, 76, 77],
    "performed_actions": [80, 81, 82],
    "planned_actions": [87, 88, 89],
    "preventive_measures": [92, 93, 94],
}


def build_export(snapshot, timezone_name="Asia/Yekaterinburg"):
    from copy import deepcopy
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = Document(TEMPLATE)
    body = list(d.element.body)
    values = _values(snapshot, timezone_name)
    overflow = []
    size = d.styles["Normal"].font.size.pt
    font = _font(round(size * 10))
    section = d.sections[0]
    width = section.page_width.mm - section.left_margin.mm - section.right_margin.mm

    def paragraph(index):
        return Paragraph(body[index], d)

    def table(index):
        return Table(body[index], d)

    def fit(key, available, count=1):
        text = str(values.get(key) or "")
        if len(_lines(text, available * 72 / 25.4 * 10, font)) <= count:
            return text
        label = FIELDS.get(key, {}).get(
            "label_ru",
            {
                "prepared_position": "Должность составителя",
                "person_0_name": "ФИО ознакомленного лица 1",
                "person_0_position": "Должность ознакомленного лица 1",
                "person_1_name": "ФИО ознакомленного лица 2",
                "person_1_position": "Должность ознакомленного лица 2",
            }.get(key, key),
        )
        overflow.append((label, text))
        return f"См. доп. {len(overflow)}"

    def insert_inline(target, text):
        if not text:
            return
        for run in target.runs:
            if "\t" not in run.text:
                continue
            before, _, after = run.text.partition("\t")
            run.text = before
            value = target.add_run(" " + text + " ")
            value.font.name, value.font.size = "Times New Roman", Pt(size)
            run._r.addnext(value._r)
            tab = target.add_run("\t" + after)
            if run._r.rPr is not None:
                tab._r.insert(0, deepcopy(run._r.rPr))
            value._r.addnext(tab._r)
            return
        raise ValueError("Incident card template inline slot has no fill tab")

    for index, key in INLINE_FIELDS.items():
        target = paragraph(index)
        label = target.text.replace("\t", " ").strip()
        available = max(18, width - font.getlength(label) / 10 * 25.4 / 72 - 7)
        insert_inline(target, fit(key, available))
    for key, indexes in NARRATIVE_FIELDS.items():
        text = fit(key, width - 4, len(indexes))
        lines = _lines(text, (width - 4) * 72 / 25.4 * 10, font) if text else []
        for index, line in zip(indexes, lines):
            paragraph(index).add_run(line)

    def tick(target, selected):
        if not selected:
            return
        for run in target.runs:
            if "☐" in run.text:
                run.text = run.text.replace("☐", "☒", 1)
                return
        raise ValueError("Incident card template checkbox slot is missing")

    for index, category in enumerate(range(1, 9), 23):
        tick(paragraph(index), category in (values.get("categories") or []))
    for index, key in [(31, "completed"), (35, "unacceptable_event")]:
        for column, choice in [(1, True), (2, False)]:
            tick(table(index).cell(0, column).paragraphs[0], values.get(key) is choice)

    def choice_grid(index, key, options, row_offset=0):
        cells = table(index)
        for row, pair in enumerate(options, row_offset):
            for column, option in enumerate(pair):
                tick(
                    cells.cell(row, column).paragraphs[0],
                    option["value"] in (values.get(key) or []),
                )

    options = FIELDS["unacceptable_event_kinds"]["options"]
    choice_grid(37, "unacceptable_event_kinds", list(zip(options[::2], options[1::2])))
    options = FIELDS["incident_types"]["options"]
    for index, start, count in [(42, 0, 4), (44, 8, 6), (46, 20, 4)]:
        choice_grid(
            index,
            "incident_types",
            list(
                zip(
                    options[start : start + count],
                    options[start + count : start + 2 * count],
                )
            ),
        )
    tick(paragraph(47), "unknown" in (values.get("incident_types") or []))
    for column, key in [(0, "actors"), (1, "motives")]:
        for row, option in enumerate(FIELDS[key]["options"], 1):
            tick(
                table(73).cell(row, column).paragraphs[0],
                option["value"] in (values.get(key) or []),
            )

    for index, name, position in [
        (97, "prepared_by", "prepared_position"),
        (99, "person_0_name", "person_0_position"),
        (101, "person_1_name", "person_1_position"),
    ]:
        for column, key in [(0, name), (1, position)]:
            target = table(index).cell(0, column).paragraphs[0]
            text = fit(key, table(index).columns[column].width.mm - 6)
            if text:
                for run in target.runs:
                    run.text = ""
                target.add_run(text).underline = True

    preparers = snapshot.get("fields", {}).get("prepared_by", {}).get("value")
    for label, people in [
        ("Составил", preparers[1:] if isinstance(preparers, list) else []),
        (
            "С результатами проведенного расследования ознакомлен(ы):",
            (values.get("acquainted") or [])[2:],
        ),
    ]:
        if people:
            overflow.append(
                (
                    label,
                    "\n".join(
                        " / ".join(
                            _clean(person.get(key) or "")
                            for key in ("name", "position")
                        )
                        for person in people
                    ),
                )
            )
    incident_id = values.get("id") or snapshot.get("meta", {}).get("case_id", "")
    for index, (label, text) in enumerate(overflow, 1):
        d.add_page_break()
        target = d.add_paragraph(
            f"Продолжение. Инцидент {incident_id}. Пункт {index}: {label}"
        )
        target.runs[0].bold = True
        target.paragraph_format.keep_with_next = True
        d.add_paragraph(text)
    buffer = BytesIO()
    d.save(buffer)
    return dict(
        filename=f"incident-card-{snapshot.get('meta', {}).get('case_id', 0)}.docx",
        revision=snapshot.get("revision", 0),
        docx_base64=base64.b64encode(buffer.getvalue()).decode("ascii"),
        html=_print_html(d, base_pages=4),
        warnings=snapshot.get("warnings", []),
    )


def _print_html(document, base_pages=4):
    """Print the native paragraphs/tables, sharing values and explicit page breaks with DOCX."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    def paragraph_html(paragraph):
        fmt = paragraph.paragraph_format
        styles = []
        for name, value in [
            ("margin-top", fmt.space_before),
            ("margin-bottom", fmt.space_after),
            ("margin-left", fmt.left_indent),
            ("text-indent", fmt.first_line_indent),
        ]:
            if value is not None:
                styles.append(f"{name}:{value.pt}pt")
        if fmt.line_spacing is not None and not isinstance(fmt.line_spacing, float):
            styles.append(f"line-height:{fmt.line_spacing.pt}pt")
        if paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            styles.append("text-align:center")
        elif paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT:
            styles.append("text-align:right")
        borders = paragraph._p.find(qn("w:pPr"))
        if borders is not None and borders.find(qn("w:pBdr")) is not None:
            styles.append("border-bottom:0.5pt solid black")
        if isinstance(fmt.line_spacing, float):
            font_size = max(
                [run.font.size.pt if run.font.size else 10.5 for run in paragraph.runs]
                or [10.5]
            )
            styles.append(
                f"line-height:{round(font_size * 1.15 * fmt.line_spacing, 2)}pt"
            )
        runs = []
        for run in paragraph.runs:
            size = f"font-size:{run.font.size.pt}pt;" if run.font.size else ""
            weight = "font-weight:bold;" if run.bold else ""
            weight += "text-decoration:underline;" if run.underline else ""
            runs.append(
                f'<span style="{size}{weight}">{escape(run.text).replace(chr(10), "<br>")}</span>'
            )
        content = "".join(runs) if paragraph.text else "&#160;"
        if "\t" in paragraph.text:
            prefix, _, suffix = paragraph.text.rpartition("\t")
            last_line = prefix.rsplit("\n", 1)[-1]
            available = max(
                3, 170 - _font(105).getlength(last_line) / 10 * 25.4 / 72 - 3
            )
            content = (
                escape(prefix).replace("\n", "<br>")
                + f'<span class="fill-line" style="width:{available}mm"></span>'
                + escape(suffix)
            )
        return f'<p style="{";".join(styles)}">{content}</p>'

    def table_html(table):
        bordered = table._tbl.tblPr.find(qn("w:tblBorders"))
        bordered = bordered is not None and any(
            b.get(qn("w:val")) == "single" for b in bordered
        )
        parts = [f'<table class="{"matrix" if bordered else "options"}"><colgroup>']
        for column in table.columns:
            parts.append(f'<col style="width:{column.width.mm}mm">')
        parts.append("</colgroup>")
        for row in table.rows:
            parts.append("<tr>")
            for cell in row.cells:
                props = cell._tc.get_or_add_tcPr()
                border = props.find(qn("w:tcBorders"))
                style = (
                    ' style="border-bottom:0.5pt solid black"'
                    if border is not None
                    else ""
                )
                parts.append(
                    f"<td{style}>"
                    + "".join(paragraph_html(p) for p in cell.paragraphs)
                    + "</td>"
                )
            parts.append("</tr>")
        parts.append("</table>")
        return "".join(parts)

    sections = []
    current = []
    for body_index, element in enumerate(document.element.body):
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            is_header = body_index in (0, 34, 59, 84)
            if paragraph.paragraph_format.page_break_before or (is_header and current):
                if current:
                    sections.append(current)
                    current = []
            if element.xpath('.//w:br[@w:type="page"]'):
                if current:
                    sections.append(current)
                    current = []
            else:
                current.append((paragraph.text, paragraph_html(paragraph)))
        elif element.tag == qn("w:tbl"):
            current.append(("", table_html(Table(element, document))))
    sections.append(current)
    html = []
    for index, section in enumerate(sections):
        continuation = index >= base_pages
        cls = "continuation" if continuation else "form-page"
        html.append(
            f'<section class="{cls}">'
            + "".join(markup for _, markup in section)
            + "</section>"
        )
    return (
        """<!doctype html><html lang="ru"><meta charset="utf-8"><title>Карточка данных об инциденте</title><style>
    @page{size:A4;margin:0}@page continuation{size:A4;margin:18mm 12mm 16mm 27mm}*{box-sizing:border-box}body{margin:0;background:#ddd;font-family:"Times New Roman",serif;font-size:10.5pt;color:black}
    p{margin:0;line-height:12.5pt;white-space:pre-wrap;overflow-wrap:anywhere}
    .form-page{width:210mm;min-height:297mm;padding:18mm 12mm 16mm 27mm;margin:0 auto;background:white;break-after:page}
    .fill-line{display:inline-block;flex:1;border-bottom:0.5pt solid black;min-width:1mm;margin-left:2mm}
    table{width:171mm;table-layout:fixed;border-collapse:collapse;font-size:inherit}td{vertical-align:top;padding:0}td p{line-height:12pt}.options td{padding-right:2mm}
    tr{break-inside:avoid}.matrix td{border:0.5pt solid black;padding:1.25pt}
    .continuation{width:210mm;margin:0 auto;padding:18mm 12mm 16mm 27mm;background:white;break-before:page}
    @media print{body{background:white}.form-page:last-child{break-after:auto}.continuation{page:continuation;width:auto;margin:0;padding:0}}
    </style><body>"""
        + "".join(html)
        + "</body></html>"
    )
